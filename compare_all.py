"""
compare_all.py - Comparaison de tous les modèles.

Lit les fichiers de métriques produits par les scripts d'entraînement et crée :
  - output/metrics_summary.xlsx  : onglet Global + un onglet par magasin
                                   + onglet Détail produits ; meilleur modèle en vert
  - output/rapport_metriques.txt : commentaire automatique
  - output/metrics_all.csv       : tableau unique (global, store, item, catégorie) avec
                                   WRMSSE, lu par generate_workbook.py et par le VBA
  - output/winner.json           : meilleur modèle JOURNALIER (RMSE global le plus bas),
                                   utilisé par predict_next_30_days.py

Le modèle mensuel (metrics_monthly.csv) est exclu des classements journaliers :
ses erreurs sont en unités / mois, non comparables aux erreurs en unités / jour.
Il a son propre onglet "Mensuel".

Métriques : MAE, RMSE (unités), MAPE (hors jours à 0, peu fiable sur les séries
intermittentes), WAPE (somme |erreur| / somme ventes) et WRMSSE (RMSE rapporté
au naïf "veille" de chaque série, pondéré par le chiffre d'affaires, comme
dans la compétition M5 ; calculé ici au niveau série).
"""
import json
import os
import sys
import traceback

import numpy as np
import pandas as pd

from utils import (CATEGORIES, OUTPUT_DIR, STORES, load_long,
                   metrics_table, rmsse_by_series, split_date, step, wrmsse)

METRIC_FILES = {
    "XGBoost passif": os.path.join(OUTPUT_DIR, "metrics_passive_ml.csv"),
    "LightGBM actif": os.path.join(OUTPUT_DIR, "metrics_active_ml.csv"),
    "Prophet (2 modes)": os.path.join(OUTPUT_DIR, "metrics_prophet.csv"),
    "LightGBM mensuel": os.path.join(OUTPUT_DIR, "metrics_monthly.csv"),
}
# Prévisions journalières : servent à recalculer les métriques par catégorie
# (les scripts d'entraînement n'écrivent pas cat_id dans leurs fichiers)
FORECAST_FILES = {
    ("XGBoost", "Passif"): os.path.join(OUTPUT_DIR, "forecasts_passive_ml.csv"),
    ("LightGBM", "Actif"): os.path.join(OUTPUT_DIR, "forecasts_active_ml.csv"),
    ("Prophet", "Passif"): os.path.join(OUTPUT_DIR, "forecasts_prophet_passive.csv"),
    ("Prophet", "Actif"): os.path.join(OUTPUT_DIR, "forecasts_prophet_active.csv"),
}
MONTHLY = "Mensuel"     # valeur de la colonne approach du modèle mensuel
SUMMARY_FILE = os.path.join(OUTPUT_DIR, "metrics_summary.xlsx")
REPORT_FILE = os.path.join(OUTPUT_DIR, "rapport_metriques.txt")
METRICS_ALL_FILE = os.path.join(OUTPUT_DIR, "metrics_all.csv")
WINNER_FILE = os.path.join(OUTPUT_DIR, "winner.json")

# Fichier à utiliser pour prédire les 30 prochains jours selon le gagnant
MODEL_FILE_MAP = {
    ("XGBoost", "Passif"): "model_passive.joblib",
    ("LightGBM", "Actif"): "model_active.joblib",
    ("Prophet", "Passif"): "forecast_prophet_future_30d.csv",
    ("Prophet", "Actif"): "forecast_prophet_future_30d.csv",
}

COLS = ["Modèle", "Approche", "MAE", "MAPE", "RMSE", "WAPE", "WRMSSE"]


def load_metrics():
    """Concatène tous les fichiers de métriques disponibles."""
    step("Lecture des fichiers de métriques...")
    frames = []
    for name, path in METRIC_FILES.items():
        if os.path.exists(path):
            frames.append(pd.read_csv(path))
            print(f"   OK      {os.path.basename(path)}")
        else:
            print(f"   ABSENT  {os.path.basename(path)} ({name} non entraîné)")
    if not frames:
        raise FileNotFoundError("Aucun fichier de métriques trouvé. Lancez d'abord les entraînements.")
    mt = pd.concat(frames, ignore_index=True)
    mt = mt.rename(columns={"model": "Modèle", "approach": "Approche"})
    mt["label"] = mt["Modèle"] + " (" + mt["Approche"] + ")"
    return mt


def category_metrics():
    """
    Métriques par catégorie des modèles journaliers, recalculées depuis les prévisions.
    Niveau "category" : store_id = magasin (ou ALL pour les 3 stores), item_id = catégorie.
    """
    step("Calcul des métriques par catégorie (depuis les fichiers de prévisions)...")
    frames = []
    for (model, approach), path in FORECAST_FILES.items():
        if not os.path.exists(path):
            continue
        fc = pd.read_csv(path)
        # La catégorie est le préfixe de l'identifiant produit (FOODS_3_090 -> FOODS)
        fc["cat_id"] = fc["item_id"].str.split("_").str[0]
        per_store = metrics_table(fc, model, approach)
        all_stores = metrics_table(fc.assign(store_id="ALL"), model, approach)
        frames += [all_stores[all_stores["level"] == "category"],
                   per_store[per_store["level"] == "category"]]
    if not frames:
        return pd.DataFrame()
    mt = pd.concat(frames, ignore_index=True)
    mt = mt.rename(columns={"model": "Modèle", "approach": "Approche"})
    mt["label"] = mt["Modèle"] + " (" + mt["Approche"] + ")"
    print(f"   {mt['label'].nunique()} modèle(s) × {mt['item_id'].nunique()} catégories")
    return mt


def wrmsse_metrics():
    """
    WRMSSE des modèles journaliers aux niveaux global, store, catégorie (ALL et par store).
    Échelle = naïf "veille" sur l'historique d'entraînement ; poids = CA des 28 derniers
    jours d'entraînement. WRMSSE < 1 : le modèle bat le naïf (pondéré par le CA).
    """
    step("Calcul du WRMSSE (métrique officielle M5, niveau série)...")
    long = load_long()
    history = long[long["date"] < split_date(long)]
    rows = []
    for (model, approach), path in FORECAST_FILES.items():
        if not os.path.exists(path):
            continue
        fc = pd.read_csv(path, parse_dates=["date"])
        per = rmsse_by_series(fc, history)
        per["cat_id"] = per["item_id"].str.split("_").str[0]
        base = {"Modèle": model, "Approche": approach}
        rows.append({**base, "level": "global", "store_id": "ALL", "item_id": "ALL", "WRMSSE": wrmsse(per)})
        for store, d in per.groupby("store_id"):
            rows.append({**base, "level": "store", "store_id": store, "item_id": "ALL", "WRMSSE": wrmsse(d)})
        for cat, d in per.groupby("cat_id"):
            rows.append({**base, "level": "category", "store_id": "ALL", "item_id": cat, "WRMSSE": wrmsse(d)})
        for (store, cat), d in per.groupby(["store_id", "cat_id"]):
            rows.append({**base, "level": "category", "store_id": store, "item_id": cat, "WRMSSE": wrmsse(d)})
    wr = pd.DataFrame(rows)
    if not wr.empty:
        g = wr[wr["level"] == "global"].sort_values("WRMSSE")
        for _, r in g.iterrows():
            print(f"   {r['Modèle']} ({r['Approche']}) : WRMSSE = {r['WRMSSE']:.3f}")
    return wr


def add_wrmsse(mt, wr):
    """Ajoute la colonne WRMSSE (jointure sur modèle / niveau / store / item)."""
    keys = ["Modèle", "Approche", "level", "store_id", "item_id"]
    if wr.empty:
        return mt.assign(WRMSSE=np.nan)
    return mt.drop(columns="WRMSSE", errors="ignore").merge(wr, on=keys, how="left")


def table_for(mt, level, store="ALL"):
    """Tableau Modèle × métriques pour un niveau donné, trié par RMSE."""
    d = mt[(mt["level"] == level) & (mt["store_id"] == store)]
    return d.reindex(columns=COLS).sort_values("RMSE").reset_index(drop=True)


def write_excel(mt, mt_cat, mt_month):
    """Écrit metrics_summary.xlsx avec mise en forme et meilleur modèle en vert."""
    step("Écriture de metrics_summary.xlsx...")
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    best_fill = PatternFill("solid", fgColor="C6EFCE")   # vert clair Excel
    best_font = Font(bold=True, color="006100")

    def style_sheet(ws, n_rows, widths, highlight_first=True):
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center")
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        if highlight_first and n_rows:
            # La ligne 2 est le meilleur modèle (tableau trié par RMSE croissant)
            for cell in ws[2]:
                cell.fill, cell.font = best_fill, best_font
        ws.freeze_panes = "A2"

    widths = {"A": 14, "B": 12, "C": 10, "D": 10, "E": 10, "F": 10, "G": 10}
    with pd.ExcelWriter(SUMMARY_FILE, engine="openpyxl") as writer:
        # Onglet récapitulatif global
        g = table_for(mt, "global")
        g.to_excel(writer, sheet_name="Global", index=False)
        style_sheet(writer.sheets["Global"], len(g), widths)
        note_row = len(g) + 3
        ws = writer.sheets["Global"]
        ws.cell(row=note_row, column=1, value="Meilleur modèle (RMSE le plus bas) surligné en vert.")
        ws.cell(row=note_row + 1, column=1, value="MAPE calculée uniquement sur les jours avec ventes > 0 "
                                                  "(peu fiable sur HOBBIES / HOUSEHOLD).")
        ws.cell(row=note_row + 2, column=1, value="WAPE = somme |erreur| / somme ventes. WRMSSE < 1 = meilleur "
                                                  "que le naïf 'veille' (pondéré par le CA, comme M5).")

        # Un onglet par magasin
        for store in STORES:
            t = table_for(mt, "store", store)
            t.to_excel(writer, sheet_name=store, index=False)
            style_sheet(writer.sheets[store], len(t), widths)

        # Détail par produit : meilleur modèle de chaque série en vert
        items = mt[mt["level"] == "item"].reindex(columns=["store_id", "item_id"] + COLS)
        items = items.sort_values(["store_id", "item_id", "RMSE"]).reset_index(drop=True)
        items.to_excel(writer, sheet_name="Détail produits", index=False)
        ws = writer.sheets["Détail produits"]
        style_sheet(ws, 0, {"A": 10, "B": 16, "C": 14, "D": 12, "E": 10, "F": 10, "G": 10},
                    highlight_first=False)
        first_of_series = ~items.duplicated(["store_id", "item_id"])
        for i in np.where(first_of_series)[0]:
            for cell in ws[int(i) + 2]:
                cell.fill, cell.font = best_fill, best_font
        ws.auto_filter.ref = ws.dimensions

        # Onglet "Par catégorie" : modèles journaliers, triés par RMSE dans chaque
        # groupe (store × catégorie, ALL = 3 stores réunis), meilleur en vert
        if not mt_cat.empty:
            write_grouped(writer, "Par catégorie", mt_cat.rename(columns={"item_id": "Catégorie"}),
                          ["store_id", "Catégorie"], best_fill, best_font, style_sheet,
                          order={"store_id": ["ALL"] + STORES, "Catégorie": CATEGORIES})

        # Onglet "Mensuel" : modèle LightGBM mensuel (+ référence naïve M-12).
        # La colonne Horizon rappelle que les erreurs sont en unités / mois.
        if not mt_month.empty:
            m = mt_month[mt_month["level"].isin(["global", "store", "category"])].copy()
            h = m["horizon"] if "horizon" in m.columns else pd.Series(1, index=m.index)
            m["Horizon"] = h.fillna(1).astype(int).astype(str) + " mois"
            if "coverage_80" in m.columns:
                m["Couverture IC 80 %"] = m["coverage_80"]
            m["Niveau"] = m["level"].map({"global": "Global", "store": "Store", "category": "Catégorie"})
            m["Catégorie"] = m["item_id"].where(m["level"] == "category", "ALL")
            m = m.reindex(columns=["Niveau", "store_id", "Catégorie", "Modèle", "Approche", "Horizon",
                                   "MAE", "MAPE", "RMSE", "WAPE", "Couverture IC 80 %"])
            ws = write_grouped(writer, "Mensuel", m, ["Niveau", "store_id", "Catégorie"],
                               best_fill, best_font, style_sheet,
                               order={"Niveau": ["Global", "Store", "Catégorie"],
                                      "store_id": ["ALL"] + STORES, "Catégorie": ["ALL"] + CATEGORIES})
            note = ws.max_row + 2
            ws.cell(row=note, column=1, value="Erreurs exprimées en unités / mois / série "
                                               "(non comparables aux modèles journaliers).")
            ws.cell(row=note + 1, column=1, value="Naïf M-12 = ventes du même mois l'année précédente "
                                                   "(référence minimale à battre).")
            ws.cell(row=note + 2, column=1, value="LightGBM M+h : modèle direct à h mois (aucune prévision "
                                                   "réinjectée). Couverture IC 80 % : part des ventes réelles "
                                                   "dans l'intervalle (objectif 80 %).")
    print(f"   -> {SUMMARY_FILE}")


def write_grouped(writer, sheet, d, keys, best_fill, best_font, style_sheet, order):
    """Écrit un tableau groupé par `keys`, trié par RMSE ; 1re ligne de chaque groupe en vert."""
    d = d.copy()
    for col, values in order.items():
        d[col] = pd.Categorical(d[col], categories=values, ordered=True)
    cols = keys + [c for c in ["Modèle", "Approche", "Horizon", "MAE", "MAPE", "RMSE", "WAPE",
                               "WRMSSE", "Couverture IC 80 %"] if c in d.columns]
    d = d.sort_values(keys + ["RMSE"])[cols].reset_index(drop=True)
    d.to_excel(writer, sheet_name=sheet, index=False)
    ws = writer.sheets[sheet]
    style_sheet(ws, 0, {c: w for c, w in zip("ABCDEFGHIJK", [11, 12, 12, 14, 10, 10, 10, 10, 10, 10, 12])},
                highlight_first=False)
    first = ~d.duplicated(keys)
    for i in np.where(first)[0]:
        for cell in ws[int(i) + 2]:
            cell.fill, cell.font = best_fill, best_font
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_metrics_all(mt, mt_cat):
    """
    Tableau unique de métriques pour Excel (feuille METRIQUES / formulaire frmMetriques).
    Colonnes : Modele, Approche, Level, Store, Item, MAE, MAPE, RMSE, WAPE, WRMSSE.
    Level = global | store | item | category (pour category, Item = catégorie).
    """
    step("Écriture de metrics_all.csv...")
    frames = [mt] + ([mt_cat] if not mt_cat.empty else [])
    d = pd.concat(frames, ignore_index=True)
    d = d.rename(columns={"Modèle": "Modele", "level": "Level", "store_id": "Store", "item_id": "Item"})
    d = d.reindex(columns=["Modele", "Approche", "Level", "Store", "Item",
                           "MAE", "MAPE", "RMSE", "WAPE", "WRMSSE"])
    order = {"global": 0, "store": 1, "category": 2, "item": 3}
    d = (d.assign(_o=d["Level"].map(order))
          .sort_values(["_o", "Store", "Item", "RMSE"]).drop(columns="_o"))
    d.to_csv(METRICS_ALL_FILE, index=False)
    print(f"   {len(d)} lignes -> {METRICS_ALL_FILE}")


def write_winner(mt):
    """
    Identifie le gagnant parmi les modèles JOURNALIERS uniquement (le mensuel, dont les
    RMSE sont en unités / mois, est déjà exclu de mt) et écrit winner.json.
    """
    step("Identification du modèle gagnant...")
    global_j = mt[mt["level"] == "global"].copy()
    global_j["label"] = global_j["Modèle"] + " (" + global_j["Approche"] + ")"
    best_row = global_j.sort_values("RMSE").iloc[0]
    model_file = MODEL_FILE_MAP.get((best_row["Modèle"], best_row["Approche"]), "model_active.joblib")

    winner = {
        "model": best_row["Modèle"],
        "approach": best_row["Approche"],
        "label": best_row["label"],
        "RMSE": float(best_row["RMSE"]),
        "MAE": float(best_row["MAE"]),
        "WAPE": float(best_row["WAPE"]),
        "model_file": model_file,
        "updated_at": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M"),
    }
    with open(WINNER_FILE, "w", encoding="utf-8") as fh:
        json.dump(winner, fh, indent=2, ensure_ascii=False)
    print(f"   -> winner.json : {winner['label']} (RMSE={winner['RMSE']:.4f}, fichier {model_file})")
    return winner


def pct_gain(old, new):
    """Amélioration relative en % (positive = new meilleur que old)."""
    return (old - new) / old * 100


def write_report(mt, mt_cat):
    """Rapport texte avec commentaires générés automatiquement."""
    step("Rédaction du rapport...")
    g = mt[mt["level"] == "global"].set_index("label")
    ranking = g.sort_values("RMSE")
    lines = ["RAPPORT DE COMPARAISON DES MODÈLES - M5 Walmart (CA_1, TX_1, WI_1)",
             "=" * 70,
             "Période de test : 180 derniers jours. Métrique principale : RMSE.",
             "MAPE calculée uniquement sur les jours avec ventes > 0 (peu fiable sur les",
             "séries intermittentes) -> WAPE et WRMSSE sont les métriques relatives de référence.", ""]

    # Classement global
    lines.append("1. CLASSEMENT GLOBAL")
    for rank, (label, r) in enumerate(ranking.iterrows(), start=1):
        lines.append(f"   {rank}. {label:<22} RMSE={r['RMSE']:.3f}  MAE={r['MAE']:.3f}  "
                     f"WAPE={r.get('WAPE', np.nan):.1f}%  WRMSSE={r.get('WRMSSE', np.nan):.3f}  "
                     f"MAPE={r['MAPE']:.1f}%")
    best = ranking.index[0]
    lines.append(f"\n   -> Modèle gagnant : {best} (RMSE {ranking.iloc[0]['RMSE']:.3f}).")
    if len(ranking) > 1:
        lines.append(f"      Il devance {ranking.index[1]} de "
                     f"{pct_gain(ranking.iloc[1]['RMSE'], ranking.iloc[0]['RMSE']):.1f} % en RMSE.")
    lines.append("")

    # Comparaisons Passif vs Actif
    lines.append("2. APPORT DES VARIABLES EXTERNES (Passif -> Actif)")
    pairs = [("ML", "XGBoost (Passif)", "LightGBM (Actif)"),
             ("Prophet", "Prophet (Passif)", "Prophet (Actif)")]
    for name, p, a in pairs:
        if p in g.index and a in g.index:
            for m in ("RMSE", "MAE"):
                gain = pct_gain(g.loc[p, m], g.loc[a, m])
                verb = "améliore" if gain > 0 else "dégrade"
                lines.append(f"   {name:<8}: l'approche active {verb} le {m} de {abs(gain):.1f} % "
                             f"({g.loc[p, m]:.3f} -> {g.loc[a, m]:.3f})")
        else:
            lines.append(f"   {name:<8}: comparaison impossible (métriques manquantes)")
    lines.append("")

    # ML vs Prophet
    ml = [l for l in g.index if l.startswith(("XGBoost", "LightGBM"))]
    pr = [l for l in g.index if l.startswith("Prophet")]
    if ml and pr:
        best_ml = g.loc[ml, "RMSE"].idxmin()
        best_pr = g.loc[pr, "RMSE"].idxmin()
        gain = pct_gain(g.loc[best_pr, "RMSE"], g.loc[best_ml, "RMSE"])
        lines.append("3. MACHINE LEARNING vs PROPHET")
        if gain > 0:
            lines.append(f"   Le meilleur ML ({best_ml}) bat le meilleur Prophet ({best_pr}) de {gain:.1f} % en RMSE.")
        else:
            lines.append(f"   Le meilleur Prophet ({best_pr}) bat le meilleur ML ({best_ml}) de {-gain:.1f} % en RMSE.")
        lines.append("")

    # Gagnant par magasin
    lines.append("4. MEILLEUR MODÈLE PAR MAGASIN")
    s = mt[mt["level"] == "store"]
    for store in STORES:
        d = s[s["store_id"] == store].sort_values("RMSE")
        if d.empty:
            continue
        lines.append(f"   {store} : {d.iloc[0]['label']:<22} RMSE={d.iloc[0]['RMSE']:.3f} "
                     f"(pire : {d.iloc[-1]['label']}, RMSE={d.iloc[-1]['RMSE']:.3f})")
    lines.append("")

    # Nombre de séries gagnées par modèle
    lines.append("5. NOMBRE DE SÉRIES (magasin × produit) GAGNÉES PAR MODÈLE")
    it = mt[mt["level"] == "item"]
    wins = it.loc[it.groupby(["store_id", "item_id"])["RMSE"].idxmin(), "label"].value_counts()
    for label, n in wins.items():
        lines.append(f"   {label:<22} {n} série(s)")
    lines.append("")

    # Gagnant par catégorie (3 stores réunis, modèles journaliers)
    if not mt_cat.empty:
        lines.append("6. MEILLEUR MODÈLE PAR CATÉGORIE")
        c = mt_cat[mt_cat["store_id"] == "ALL"]
        for cat in CATEGORIES:
            d = c[c["item_id"] == cat].sort_values("RMSE")
            if d.empty:
                continue
            lines.append(f"   {cat:<10}: {d.iloc[0]['label']:<22} RMSE={d.iloc[0]['RMSE']:.3f}  "
                         f"WAPE={d.iloc[0]['WAPE']:.1f}%  WRMSSE={d.iloc[0].get('WRMSSE', np.nan):.3f}")
        lines.append("")

    # WRMSSE : lecture de la métrique officielle M5
    if "WRMSSE" in g.columns and g["WRMSSE"].notna().any():
        lines.append("7. WRMSSE (métrique officielle M5, niveau série)")
        for label, r in g.sort_values("WRMSSE").iterrows():
            verdict = "bat" if r["WRMSSE"] < 1 else "ne bat pas"
            lines.append(f"   {label:<22} WRMSSE={r['WRMSSE']:.3f} -> {verdict} le naïf 'veille' "
                         f"(pondéré par le CA)")
        lines.append("   Rappel : le gagnant du concours M5 obtenait ~0.52 sur les 12 niveaux "
                     "d'agrégation (ici : niveau série seul, horizon 180 j -> non comparable directement).")

    text = "\n".join(lines)
    with open(REPORT_FILE, "w", encoding="utf-8") as f:
        f.write(text + "\n")
    print(text)
    print(f"\n   -> {REPORT_FILE}")


def main():
    print("=" * 70)
    print(" COMPARAISON DE TOUS LES MODÈLES")
    print("=" * 70)
    mt_all = load_metrics()
    # Séparation journalier / mensuel (unités différentes : jour vs mois)
    mt = mt_all[mt_all["Approche"] != MONTHLY]
    mt_month = mt_all[mt_all["Approche"] == MONTHLY]
    if mt.empty:
        raise FileNotFoundError("Aucune métrique journalière trouvée. Lancez d'abord les entraînements.")
    mt_cat = category_metrics()
    # WRMSSE ajouté à toutes les tables journalières
    wr = wrmsse_metrics()
    mt = add_wrmsse(mt, wr)
    if not mt_cat.empty:
        mt_cat = add_wrmsse(mt_cat, wr)
    write_excel(mt, mt_cat, mt_month)
    write_report(mt, mt_cat)
    write_metrics_all(mt, mt_cat)
    write_winner(mt)
    print("\nComparaison terminée avec succès.")


if __name__ == "__main__":
    try:
        main()
    except PermissionError as e:
        print(f"\n[ERREUR] Fichier verrouillé ({e.filename}). Fermez-le dans Excel puis relancez.")
        sys.exit(1)
    except Exception as e:
        print(f"\n[ERREUR] compare_all : {e}")
        traceback.print_exc()
        sys.exit(1)

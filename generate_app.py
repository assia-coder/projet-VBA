"""
generate_app.py - Application Excel/VBA de prevision de la demande (M5 Forecasting)
Version 3.1

Genere Application_Prevision.xlsx (base "admin", sans macros) :
  - Accueil : logo dans la banniere, barre de connexion au-dessus des KPI,
    KPI 100 % formules dependants de la session (role / magasin)
  - Feuilles de donnees T_* : UNIQUEMENT des tableaux Excel (ListObjects)
  - Colonnes CALCULEES par formules Excel (aucune saisie, aucun import) :
        T_Produits.prix_moyen      <- moyenne des sell_price de T_Transactions
        T_Prix.sell_price          <- moyenne hebdomadaire des sell_price de T_Transactions
        T_Calendrier.weekday/month/year <- calcules a partir de la date
  - Gestion : liste des tables + boutons (poses par build_userforms.py)
  - Bouton « Accueil » sur toutes les feuilles
  - R_Session (cachee) : role / magasin / utilisateur lus par les formules

Sources :
  output/series_long.csv                  -> T_Transactions (toutes les colonnes, toutes les lignes)
  output/top_produits_par_categorie.csv   -> T_Produits
  data/calendar.csv                       -> T_Calendrier
  deduit                                  -> T_Categories, T_Magasins, T_Prix (cles semaine x magasin x produit)

Usage :
    python generate_app.py
    python build_userforms.py      (Windows + Excel)
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

# ------------------------------------------------------------------ Chemins
ROOT = Path(__file__).parent
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "output"
VBA_DIR = ROOT / "vba"
XLSX_PATH = ROOT / "Application_Prevision.xlsx"
LOGO_PATH = ROOT / "Walmart-Logo.jpg"
SEARCH_DIRS = [OUTPUT_DIR, DATA_DIR, ROOT]

# ------------------------------------------------------------------ Perimetre
SCOPE_STORES = ["CA_1", "TX_1", "WI_1"]
# None = on garde tout le contenu de top_produits_par_categorie.csv.
# Mettre 5 pour ne garder que les 5 produits les plus vendus PAR MAGASIN.
TOP_N_PER_STORE = None

# Colonne des quantites vendues dans series_long.csv (None = detection automatique)
SALES_COL = None

STORE_INFO = {  # store_id -> (state_id, description)
    "CA_1": ("CA", "California"),
    "TX_1": ("TX", "Texas"),
    "WI_1": ("WI", "Wisconsin"),
}
CAT_INFO = {
    "FOODS": "Produits alimentaires",
    "HOBBIES": "Loisirs et passe-temps",
    "HOUSEHOLD": "Articles pour la maison",
}

# ------------------------------------------------------------------ Palette Walmart
P = {
    "primary": "0071CE", "dark": "004F9E", "accent": "FFC220", "light": "78B9E7",
    "very_light": "E6F1FA", "white": "FFFFFF", "bg": "F5F7FA", "border": "D0D5DD",
    "text": "1A1A1A", "muted": "667085", "success": "12B76A", "warning": "F79009",
    "error": "D92D20",
}
PASSWORD = "walmart2024"

# ------------------------------------------------------------------ Layout des feuilles de donnees
HEADER_ROW = 6      # en-tete du tableau (rows 2-4 : bandeau, 5 : espace)
DATA_START_ROW = 7  # 1re ligne de donnees

DATA_SHEETS = ["T_Produits", "T_Categories", "T_Magasins",
               "T_Transactions", "T_Prix", "T_Calendrier"]
TABLE_NAMES = {
    "T_Produits": "tblProduits", "T_Categories": "tblCategories",
    "T_Magasins": "tblMagasins", "T_Transactions": "tblTransactions",
    "T_Prix": "tblPrix", "T_Calendrier": "tblCalendrier",
}
COLS = {}  # nom de feuille -> liste des en-tetes (rempli a la generation)

# ------------------------------------------------------------------ Layout Accueil
# Les boutons sont des formes dessinees par build_userforms.py (avec une marge interne
# -> espace visible entre deux boutons voisins). Les cellules sous les boutons restent vides.
# (ancre, libelle, macro VBA, couleur de fond, couleur de texte)
ACCUEIL_ROW_CONNEXION = 6
ACCUEIL_BUTTONS = [
    ("H6:I6", "▶  Se connecter", "OpenLogin", P["primary"], P["white"]),
    ("J6:K6", "✕  Se déconnecter", "Logout", P["error"], P["white"]),
    ("B13:C13", "▶  Dashboard transactionnel", "OpenPowerBI", P["primary"], P["white"]),
    ("D13:E13", "▶  Résultats Machine Learning", "OpenML", P["primary"], P["white"]),
    ("F13:G13", "▶  Prévision de la demande", "OpenPrevision", P["primary"], P["white"]),
    ("H13:K13", "▶  Gestion des données", "OpenGestion", P["accent"], P["dark"]),
]

# ------------------------------------------------------------------ Layout Gestion
# (libelle, description, feuille, formulaire, macro) ; formulaire/macro = None -> pas de saisie
GESTION_TABLES = [
    ("Produits", "Produits suivis (clé magasin × produit)", "T_Produits", "frmProduits", "OpenProduits"),
    ("Catégories", "Catégories de produits", "T_Categories", "frmCategories", "OpenCategories"),
    ("Magasins", "Magasins du périmètre", "T_Magasins", "frmMagasins", "OpenMagasins"),
    ("Transactions", "Historique des ventes journalières", "T_Transactions", "frmTransactions", "OpenTransactions"),
    ("Prix", "Prix hebdomadaires (calculés par formules)", "T_Prix", None, None),
    ("Calendrier", "Calendrier + événements", "T_Calendrier", "frmCalendrier", "OpenCalendrier"),
]
GESTION_FIRST_TABLE_ROW = 13
GESTION_BUTTONS = [
    ("J7:K8", "✕  Se déconnecter", "Logout", P["error"], P["white"]),
] + [
    (f"H{GESTION_FIRST_TABLE_ROW + i}:I{GESTION_FIRST_TABLE_ROW + i}", f"▶ {t[0]}", t[4],
     P["accent"], P["dark"])
    for i, t in enumerate(GESTION_TABLES) if t[4]
]

# ------------------------------------------------------------------ Bouton « Accueil » (toutes les feuilles)
HOME_BUTTON = ("←  Accueil", "GoHome", P["accent"], P["dark"])
HOME_ANCHORS = {   # feuille -> zone du bouton (coin haut-droit de la banniere)
    "PowerBI": "J2:K3", "ML": "K2:L3", "Prevision": "N2:O3", "Gestion": "J2:K3",
    "T_Produits": "J2:K3", "T_Categories": "J2:K3", "T_Magasins": "J2:K3",
    "T_Transactions": "J2:K3", "T_Prix": "J2:K3", "T_Calendrier": "J2:K3",
    "Config": "E1:F2",
    "R_Previsions": "M1:N2", "R_Metriques": "M1:N2", "R_Resume": "M1:N2",
    "R_Mensuel": "M1:N2", "R_MetMensuel": "M1:N2",
}
BUTTONS_BY_SHEET = {"Accueil": list(ACCUEIL_BUTTONS), "Gestion": list(GESTION_BUTTONS)}
for _sh, _anchor in HOME_ANCHORS.items():
    BUTTONS_BY_SHEET.setdefault(_sh, []).append((_anchor,) + HOME_BUTTON)

# Largeurs de colonnes des tableaux (defaut : 14)
COL_WIDTHS = {
    "store_id": 11, "item_id": 24, "cat_id": 13, "dept_id": 13, "date": 12,
    "description": 26, "event_name": 22, "event_type": 14, "event_name_1": 24,
    "event_name_2": 24, "event_type_1": 14, "event_type_2": 14, "source": 10,
    "saisi_le": 18, "date_ajout": 12, "weekday": 12, "month": 9, "year": 9,
    "prix_moyen": 13, "sell_price": 13, "zero_ratio": 12, "days_to_monthend": 16,
}


# ------------------------------------------------------------------ Helpers styles
def fill(c):
    return PatternFill("solid", fgColor=c)


def fnt(size=11, bold=False, color=None, italic=False):
    return Font(name="Calibri", size=size, bold=bold, italic=italic, color=color or P["text"])


def bord(color=None, style="thin"):
    s = Side(style=style, color=color or P["border"])
    return Border(left=s, right=s, top=s, bottom=s)


def title_banner(ws, rng, text, subtitle=None, bg=None):
    bg = bg or P["primary"]
    ws.merge_cells(rng)
    for row in ws[rng]:
        for c in row:
            c.fill = fill(bg)
    top = ws[rng.split(":")[0]]
    top.value = text
    top.font = fnt(18, bold=True, color=P["white"])
    top.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    if subtitle:
        parts = rng.split(":")
        col_l, row_l = parts[0][0], int(parts[0][1:]) + 1
        col_r, row_r = parts[1][0], int(parts[1][1:]) + 1
        rng2 = f"{col_l}{row_l}:{col_r}{row_r}"
        ws.merge_cells(rng2)
        for row in ws[rng2]:
            for c in row:
                c.fill = fill(bg)
        c = ws[f"{col_l}{row_l}"]
        c.value = subtitle
        c.font = fnt(10, italic=True, color=P["white"])
        c.alignment = Alignment(horizontal="left", vertical="center", indent=2)


def section(ws, rng, text, bg=None):
    bg = bg or P["dark"]
    ws.merge_cells(rng)
    for row in ws[rng]:
        for c in row:
            c.fill = fill(bg)
    c = ws[rng.split(":")[0]]
    c.value = text
    c.font = fnt(12, bold=True, color=P["white"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def header_row(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=lab)
        cell.font = fnt(11, bold=True, color=P["white"])
        cell.fill = fill(P["primary"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = bord(P["primary"])


def style_block(ws, rng, bg, border_color=None):
    for row in ws[rng]:
        for c in row:
            c.fill = fill(bg)
            if border_color:
                c.border = bord(border_color)


# ------------------------------------------------------------------ Helpers donnees
def find_file(*names):
    for d in SEARCH_DIRS:
        for n in names:
            p = d / n
            if p.exists():
                return p
    return None


def read_csv_required(label, *names):
    p = find_file(*names)
    if p is None:
        print(f"[ERREUR] Fichier introuvable : {names[0]} (cherché dans {[str(d) for d in SEARCH_DIRS]})")
        sys.exit(1)
    print(f"   lecture {p.relative_to(ROOT) if p.is_relative_to(ROOT) else p}")
    return pd.read_csv(p)


def clean(df):
    """DataFrame -> objets Python natifs, NaN/NaT -> None."""
    return df.astype(object).where(df.notna(), None)


def L(sheet, header):
    """Lettre de colonne Excel d'un en-tete de tableau (feuille deja ecrite)."""
    return get_column_letter(2 + COLS[sheet].index(header))


def CL(cols, header):
    """Lettre de colonne Excel d'un en-tete, pour une liste de colonnes pas encore ecrite."""
    return get_column_letter(2 + cols.index(header))


def rng(sheet, header, last=5000):
    c = L(sheet, header)
    return f"{sheet}!${c}${DATA_START_ROW}:${c}${last}"


# ------------------------------------------------------------------ Feuilles de donnees (tableaux)
def write_table_sheet(wb, sheet, title, subtitle, df, formats=None):
    ws = wb.create_sheet(sheet)
    cols = [str(c) for c in df.columns]
    COLS[sheet] = cols
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = P["primary"]
    ws.column_dimensions["A"].width = 3
    for i, c in enumerate(cols):
        ws.column_dimensions[get_column_letter(2 + i)].width = COL_WIDTHS.get(c, 14)

    ws.row_dimensions[2].height = 34
    title_banner(ws, "B2:K2", title, subtitle=subtitle)
    ws.row_dimensions[4].height = 6
    for c in ws["B4:K4"][0]:
        c.fill = fill(P["accent"])

    header_row(ws, HEADER_ROW, cols, start_col=2)
    ws.row_dimensions[HEADER_ROW].height = 28

    # Donnees : ws.append continue sous la ligne d'en-tete (colonne A laissee vide)
    n = 0
    for row in clean(df).itertuples(index=False, name=None):
        ws.append([None] + list(row))
        n += 1
    if n == 0:  # un tableau Excel doit contenir au moins 1 ligne de donnees
        ws.append([None] * (len(cols) + 1))
    last_row = HEADER_ROW + max(n, 1)

    # formats numeriques des colonnes calculees
    for col_name, fmt in (formats or {}).items():
        idx = 2 + cols.index(col_name)
        for r in range(DATA_START_ROW, last_row + 1):
            ws.cell(row=r, column=idx).number_format = fmt

    ref = f"B{HEADER_ROW}:{get_column_letter(1 + len(cols))}{last_row}"
    tab = Table(displayName=TABLE_NAMES[sheet], ref=ref,
                tableColumns=[TableColumn(id=i + 1, name=c) for i, c in enumerate(cols)])
    tab.autoFilter = None   # pas de flèches de filtre : filtrage = role, applique par VBA
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = f"A{DATA_START_ROW}"
    print(f"→ {sheet:<15} {n:>8} lignes × {len(cols)} colonnes")
    return n


def load_produits():
    df = read_csv_required("produits", "top_produits_par_categorie.csv", "top_produits_par_catégorie.csv")
    # rang / ventes_totales ne servent qu'a trier / limiter ; elles ne sont PAS ecrites dans T_Produits
    wanted = ["store_id", "item_id", "cat_id", "dept_id", "rang", "ventes_totales", "zero_ratio"]
    for c in wanted:
        if c not in df.columns:
            print(f"[warn] colonne absente de top_produits : {c}")
            df[c] = None
    df = df[wanted].copy()
    before = len(df)
    df = df[df["store_id"].isin(SCOPE_STORES)]
    if len(df) != before:
        print(f"   [info] {before - len(df)} ligne(s) hors périmètre {SCOPE_STORES} ignorée(s)")
    if TOP_N_PER_STORE:
        df = (df.sort_values("ventes_totales", ascending=False)
                .groupby("store_id", group_keys=False).head(TOP_N_PER_STORE))
    df = df.drop_duplicates(["store_id", "item_id"]).sort_values(["store_id", "cat_id", "rang"])
    df = df[["store_id", "item_id", "cat_id", "dept_id", "zero_ratio"]].copy()
    df["prix_moyen"] = None          # remplace par une formule dans main()
    df["date_ajout"] = None
    df["source"] = "M5"
    df = df[["store_id", "item_id", "cat_id", "dept_id", "prix_moyen", "zero_ratio", "date_ajout", "source"]]
    return df.reset_index(drop=True)


def detect_sales_col(df):
    if SALES_COL:
        return SALES_COL
    for c in ("sales", "demand", "unit_sales", "units", "qty", "quantity", "ventes", "y"):
        if c in df.columns:
            return c
    num = [c for c in df.select_dtypes("number").columns
           if c not in ("sell_price", "wm_yr_wk", "wday", "month", "year")]
    return num[0] if num else None


def load_transactions(produits):
    global SALES_COL
    df = read_csv_required("series_long", "series_long.csv")
    for c in ("date", "store_id", "item_id", "sell_price"):
        if c not in df.columns:
            print(f"[ERREUR] series_long.csv : colonne obligatoire absente : {c}")
            sys.exit(1)
    SALES_COL = detect_sales_col(df)
    if SALES_COL is None or SALES_COL not in df.columns:
        print("[ERREUR] Colonne des quantités vendues introuvable dans series_long.csv : "
              "renseignez SALES_COL en haut de generate_app.py.")
        sys.exit(1)
    print(f"   colonne des ventes : {SALES_COL}")
    before = len(df)
    pairs = set(zip(produits["store_id"], produits["item_id"]))
    mask = pd.MultiIndex.from_frame(df[["store_id", "item_id"]]).isin(pairs)
    df = df[mask].copy()
    if len(df) != before:
        print(f"   [info] {before - len(df)} transaction(s) hors périmètre (magasin × produit) ignorée(s)")
    df["date"] = pd.to_datetime(df["date"]).dt.date
    df["source"] = "M5"
    df["saisi_le"] = None
    return df.reset_index(drop=True)


def load_calendar():
    df = read_csv_required("calendar", "calendar.csv")
    df["date"] = pd.to_datetime(df["date"])
    return df.sort_values("date").reset_index(drop=True)


def build_prix_keys(trans, calendar):
    """Cles (semaine x magasin x produit) des prix hebdomadaires ; sell_price = formule Excel.

    date = 1er jour de la semaine M5 (samedi). Seules les semaines ou le produit a des
    transactions avec un prix sont creees (dates passees)."""
    s = trans[["date", "store_id", "item_id", "sell_price"]].copy()
    s["date"] = pd.to_datetime(s["date"])
    s = s.dropna(subset=["sell_price"])
    if "wm_yr_wk" in calendar.columns:
        wk = calendar.groupby("wm_yr_wk")["date"].min().rename("week_start").reset_index()
        s = (s.merge(calendar[["date", "wm_yr_wk"]], on="date", how="left")
              .merge(wk, on="wm_yr_wk", how="left"))
    else:   # semaine M5 : du samedi au vendredi
        s["week_start"] = s["date"] - pd.to_timedelta((s["date"].dt.dayofweek - 5) % 7, unit="D")
    g = (s[["week_start", "store_id", "item_id"]].drop_duplicates()
          .rename(columns={"week_start": "date"})
          .sort_values(["store_id", "item_id", "date"]).reset_index(drop=True))
    g["date"] = g["date"].dt.date
    return g


# ------------------------------------------------------------------ Gestion
def build_gestion(wb):
    print("→ Gestion")
    ws = wb.create_sheet("Gestion")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = P["accent"]
    for col, w in zip("ABCDEFGHIJK", [3, 30, 30, 14, 16, 16, 20, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w

    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 34
    title_banner(ws, "B2:K2", "Gestion des données",
                 subtitle="Tables de l'application : consultation, recherche, saisie contrôlée via UserForms.")
    ws.row_dimensions[4].height = 10          # espace sous la banniere

    # ---- Session utilisateur (boutons a droite, avec marge : rien n'est colle aux cellules)
    ws.row_dimensions[5].height = 22
    section(ws, "B5:K5", "  Session utilisateur", bg=P["dark"])
    ws.row_dimensions[6].height = 10
    for r in (7, 8, 9):
        ws.row_dimensions[r].height = 24
    for r, label, formula in [
        (7, "Statut session :", '=IF(sessRole="","non connecté",sessRoleLabel)'),
        (8, "Utilisateur :", '=IF(sessRole="","—",sessUser)'),
        (9, "Magasin :", '=IF(sessRole="","—",IF(sessStore="ALL","Tous les magasins",sessStore))'),
    ]:
        a = ws.cell(row=r, column=2, value=label)
        a.font = fnt(11, bold=True, color=P["dark"])
        a.alignment = Alignment(vertical="center", indent=1)
        c = ws.cell(row=r, column=3, value=formula)
        c.font = fnt(11, color=P["muted"])
        c.alignment = Alignment(vertical="center")
    # bouton « Se deconnecter » : J7:K8 (forme posee par build_userforms.py)

    # ---- Tables
    ws.row_dimensions[11].height = 22
    section(ws, "B11:K11", "  Tables de l'application", bg=P["primary"])
    header_row(ws, 12, ["Table", "Contenu", "Lignes", "dont MANUEL", "Accès", "", "Formulaire", ""], start_col=2)
    ws.merge_cells("F12:G12")
    ws.merge_cells("H12:I12")
    ws.row_dimensions[12].height = 24

    for i, (nom, desc, sheet, form, _macro) in enumerate(GESTION_TABLES):
        r = GESTION_FIRST_TABLE_ROW + i
        ws.row_dimensions[r].height = 26
        key = L(sheet, COLS[sheet][0])
        ws.cell(row=r, column=2, value=nom).font = fnt(11, bold=True, color=P["dark"])
        ws.cell(row=r, column=3, value=desc).font = fnt(10)
        # SUBTOTAL(103) ne compte que les lignes visibles : suit le filtre du role connecte
        ws.cell(row=r, column=4,
                value=f"=SUBTOTAL(103,{sheet}!${key}${DATA_START_ROW}:${key}$1048576)").font = fnt(10)
        if "source" not in COLS[sheet]:
            f_manuel = "—"
        else:
            src = L(sheet, "source")
            if "store_id" in COLS[sheet] and sheet != "T_Categories":
                st = L(sheet, "store_id")
                f_manuel = (f'=COUNTIFS({sheet}!${src}${DATA_START_ROW}:${src}$1048576,"MANUEL",'
                            f'{sheet}!${st}${DATA_START_ROW}:${st}$1048576,sessStoreCrit)')
            else:
                f_manuel = f'=COUNTIF({sheet}!${src}${DATA_START_ROW}:${src}$1048576,"MANUEL")'
        ws.cell(row=r, column=5, value=f_manuel).font = fnt(10)
        ws.merge_cells(f"F{r}:G{r}")
        if form is None:
            acces = "Lecture seule (calculé)"
        else:
            acces = '=IF(OR(sessRole="admin",sessRole="hq"),"Lecture / écriture","Lecture seule")'
        c = ws.cell(row=r, column=6, value=acces)
        c.font = fnt(10, italic=True, color=P["muted"])
        ws.merge_cells(f"H{r}:I{r}")
        if form is None:   # pas de bouton : table entierement calculee
            c = ws.cell(row=r, column=8, value="— aucune saisie —")
            c.font = fnt(10, italic=True, color=P["muted"])
            c.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(2, 8):
            ws.cell(row=r, column=col).alignment = Alignment(
                vertical="center", horizontal=("center" if col in (4, 5) else None), indent=(1 if col in (2, 3) else 0))

    section(ws, "B21:K21", "  Règles de cohérence", bg=P["dark"])
    regles = [
        "• Transactions et calendrier : lignes source = M5 en lecture seule ; les saisies utilisateur sont en source = MANUEL.",
        "• Prix moyen d'un produit et prix hebdomadaires : calculés par formules Excel à partir de la feuille Transactions (aucune saisie).",
        "• Weekday, month et year du calendrier : calculés par formules à partir de la date (non saisis).",
        "• Un produit lié à des transactions ne peut pas être supprimé. Une catégorie ou un magasin ne se supprime que s'il n'a plus de produit.",
        "• Gérant de magasin : son magasin uniquement · Analyste HQ : les 3 magasins · Administrateur : accès complet.",
        "• Les modifications manuelles n'alimentent pas les modèles Python (qui lisent output/series_long.csv).",
    ]
    for i, txt in enumerate(regles, start=22):
        ws.merge_cells(f"B{i}:K{i}")
        c = ws.cell(row=i, column=2, value=txt)
        c.font = fnt(10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 28


# ------------------------------------------------------------------ Feuilles annexes
def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names[name] = dn
    except TypeError:
        wb.defined_names.append(dn)


def build_session_sheet(wb):
    ws = wb.create_sheet("R_Session")
    ws["A1"], ws["B1"] = "key", "value"
    for i, k in enumerate(["role", "store", "username", "login_time", "role_label", "store_crit"], start=2):
        ws.cell(row=i, column=1, value=k)
    ws["B6"] = ('=IF(B2="admin","Administrateur",IF(B2="hq","Analyste HQ",'
                'IF(B2="store","Gérant de magasin","")))')
    ws["B7"] = '=IF(OR(B3="",B3="ALL"),"*",B3)'   # critere COUNTIFS : "*" = tous les magasins
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    ws.sheet_state = "hidden"
    for n, ref in [("sessRole", "R_Session!$B$2"), ("sessStore", "R_Session!$B$3"),
                   ("sessUser", "R_Session!$B$4"), ("sessRoleLabel", "R_Session!$B$6"),
                   ("sessStoreCrit", "R_Session!$B$7")]:
        add_name(wb, n, ref)


def build_powerbi(wb):
    ws = wb.create_sheet("PowerBI")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = P["primary"]
    ws.row_dimensions[2].height = 30
    title_banner(ws, "B2:K2", "PowerBI", subtitle="Dashboard transactionnel — Power BI")
    ws["B5"] = "État :"
    ws["C5"] = '=IF(cfgPowerBI="","Aucun dashboard configuré (Config!C5).","Dashboard configuré.")'
    ws["B6"] = "Lien :"
    ws["C6"] = '=IF(cfgPowerBI="","",HYPERLINK(cfgPowerBI,"▶ Ouvrir le dashboard Power BI"))'
    ws["B8"] = "Procédures :"
    for i, t in enumerate([
        "A · Publier le .pbix dans Power BI, copier l'URL dans Config!C5.",
        "B · Insertion > Compléments Office > Microsoft Power BI (zone C20:L45).",
        "C · Coller le chemin local du .pbix dans Config!C5.",
    ], start=9):
        ws.merge_cells(f"B{i}:K{i}")
        ws.cell(row=i, column=2, value=t).font = fnt(10)


def build_ml(wb):
    ws = wb.create_sheet("ML")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = P["primary"]
    ws.row_dimensions[2].height = 30
    title_banner(ws, "B2:L2", "ML", subtitle="Résultats Machine Learning — comparaison des modèles")


def build_prevision(wb):
    ws = wb.create_sheet("Prevision")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = P["primary"]
    ws.row_dimensions[2].height = 30
    title_banner(ws, "B2:O2", "Prevision", subtitle="Prévision de la demande — choisir les paramètres")


def build_config(wb):
    ws = wb.create_sheet("Config")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 50
    ws["B1"] = "Configuration"
    ws["B1"].font = fnt(14, bold=True, color=P["dark"])
    ws["B3"] = "PowerBI URL"
    ws["C3"] = ""
    ws["B4"] = "SHEET_PASSWORD"
    ws["C4"] = PASSWORD
    ws["B5"] = "pbi_path"
    ws["C5"] = ""
    add_name(wb, "cfgPowerBI", "Config!$C$5")


def build_r_sheets(wb):
    specs = {
        "R_Previsions": ["date", "store_id", "item_id", "cat_id", "model", "set", "pred", "pi_low", "pi_high"],
        "R_Metriques": ["key", "model", "level", "store", "item", "n", "MAE", "MAPE", "RMSE", "WAPE", "WRMSSE"],
        "R_Resume": ["key", "value"],
        "R_Mensuel": ["date", "horizon", "pred", "store_id", "item_id", "cat_id", "pi_low", "pi_high", ""],
        "R_MetMensuel": ["key", "model", "level", "store", "item", "MAE", "RMSE", "MAPE", "WAPE", "IC80", "n"],
    }
    for name, headers in specs.items():
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = P["muted"]
        for j, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = fnt(10, bold=True, color=P["white"])
            c.fill = fill(P["muted"])


# ------------------------------------------------------------------ Accueil
def place_logo(ws, anchor_col_idx, anchor_row_idx, box_w, box_h):
    """Logo centre dans un bloc (en pixels) ancre a (colonne, ligne) 0-based."""
    img = XLImage(str(LOGO_PATH))
    ratio = img.width / img.height if img.height else 1
    h = box_h - 24
    w = int(h * ratio)
    if w > box_w - 30:
        w = box_w - 30
        h = int(w / ratio)
    img.width, img.height = w, h
    marker = AnchorMarker(col=anchor_col_idx, row=anchor_row_idx,
                          colOff=pixels_to_EMU(max((box_w - w) // 2, 0)),
                          rowOff=pixels_to_EMU(max((box_h - h) // 2, 0)))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
    ws.add_image(img)


def build_accueil(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = P["primary"]
    for col, w in zip("ABCDEFGHIJK", [3] + [22] * 10):
        ws.column_dimensions[col].width = w

    # ---- Banniere (titre a gauche, logo a droite, MEME ligne)
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 42
    ws.row_dimensions[3].height = 24
    title_banner(ws, "B2:I2", "Prévision de la demande — M5 Forecasting",
                 subtitle="Walmart · magasins CA_1 · TX_1 · WI_1 · Excel / VBA · Python · Power BI")
    ws.merge_cells("J2:K3")
    style_block(ws, "J2:K3", P["white"], P["primary"])
    if LOGO_PATH.exists():
        # bloc J2:K3 : 2 colonnes de 22 car. (~159 px) x (42 + 24 pt = 66 pt ~ 88 px)
        place_logo(ws, anchor_col_idx=9, anchor_row_idx=1, box_w=318, box_h=88)
    else:
        c = ws["J2"]
        c.value = "Walmart"
        c.font = fnt(20, bold=True, color=P["primary"])
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 6
    for c in ws["B4:K4"][0]:
        c.fill = fill(P["accent"])
    ws.row_dimensions[5].height = 16          # espace sous la banniere

    # ---- Barre de connexion (au-dessus des KPI) : statut a gauche, boutons a droite
    ws.row_dimensions[ACCUEIL_ROW_CONNEXION].height = 36
    ws.merge_cells("B6:G6")
    style_block(ws, "B6:G6", P["very_light"], P["border"])
    c = ws["B6"]
    c.value = ('=IF(sessRole="","Non connecté — cliquez sur « Se connecter »",'
               '"Connecté : "&sessUser&"  ·  "&sessRoleLabel&"  ·  Magasin : "&IF(sessStore="ALL","Tous",sessStore))')
    c.font = fnt(11, bold=True, color=P["dark"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    # boutons H6:I6 et J6:K6 : formes posees par build_userforms.py (avec marge entre elles)
    ws.row_dimensions[7].height = 16          # espace avant les KPI

    # ---- KPI : formules dependantes de la session (role / magasin) -> toujours exacts
    ws.row_dimensions[8].height = 26
    ws.row_dimensions[9].height = 44
    ws.row_dimensions[10].height = 16
    p_store = rng("T_Produits", "store_id")
    p_cat = rng("T_Produits", "cat_id")
    t_store, t_sales, t_price = ("tblTransactions[store_id]", f"tblTransactions[{SALES_COL}]",
                                 "tblTransactions[sell_price]")
    kpis = [
        ("Produits suivis",
         f'=IF(sessRole="","—",COUNTIFS({p_store},sessStoreCrit))', "#,##0"),
        ("Magasins",
         f'=IF(sessRole="","—",COUNTIFS({rng("T_Magasins", "store_id", 200)},sessStoreCrit))', "#,##0"),
        ("Catégories",
         f'=IF(sessRole="","—",SUMPRODUCT(--(COUNTIFS({p_store},sessStoreCrit,'
         f'{p_cat},{rng("T_Categories", "cat_id", 200)})>0)))', "#,##0"),
        # KPI remplacant « Modele gagnant » : chiffre d'affaires cumule (unites x prix de vente)
        ("Chiffre d'affaires cumulé ($)",
         f'=IF(sessRole="","—",IF(sessStore="ALL",SUMPRODUCT({t_sales},{t_price}),'
         f'SUMPRODUCT(({t_store}=sessStore)*{t_sales}*{t_price})))', '"$"#,##0'),
    ]
    spans = [("B", "D"), ("E", "G"), ("H", "I"), ("J", "K")]
    for (label, formula, fmt), (c1, c2) in zip(kpis, spans):
        ws.merge_cells(f"{c1}8:{c2}8")
        ws.merge_cells(f"{c1}9:{c2}9")
        lc = ws[f"{c1}8"]
        lc.value = label
        lc.font = fnt(10, bold=True, color=P["white"])
        lc.fill = fill(P["primary"])
        lc.alignment = Alignment(horizontal="center", vertical="center")
        vc = ws[f"{c1}9"]
        vc.value = formula
        vc.number_format = fmt
        vc.font = fnt(18, bold=True, color=P["dark"])
        vc.fill = fill(P["light"])
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws[f"{c1}8:{c2}9"]:
            for c in row:
                c.border = bord(P["primary"])

    # ---- Navigation (boutons poses par build_userforms.py, sans hyperliens :
    #      les feuilles sont masquees avant la connexion)
    ws.row_dimensions[11].height = 24
    ws.merge_cells("B11:K11")
    ws["B11"].value = "  Navigation"
    ws["B11"].font = fnt(13, bold=True, color=P["white"])
    ws["B11"].fill = fill(P["dark"])
    ws["B11"].alignment = Alignment(vertical="center")
    ws.row_dimensions[12].height = 8
    ws.row_dimensions[13].height = 56
    ws.row_dimensions[14].height = 16

    # ---- Contexte
    ws.row_dimensions[15].height = 22
    ws.merge_cells("B15:K15")
    ws["B15"].value = "  Contexte et objectif"
    ws["B15"].font = fnt(12, bold=True, color=P["white"])
    ws["B15"].fill = fill(P["dark"])
    ws["B15"].alignment = Alignment(vertical="center")
    ws.merge_cells("B16:K18")
    ws["B16"] = ("Projet académique de prévision de la demande sur le jeu de données M5 Forecasting (Walmart). "
                 "L'application réunit prévision (modèles entraînés en Python), analyse des ventes (Power BI), "
                 "comparaison des modèles et gestion des tables de référence, dans un seul classeur Excel piloté par VBA.")
    ws["B16"].font = fnt(10)
    ws["B16"].alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws["B16:K18"]:
        for c in row:
            c.fill = fill(P["bg"])
            c.border = bord(P["border"])


# ------------------------------------------------------------------ MAIN
def main():
    print("=== Génération Application_Prevision (v3.1) ===")
    wb = Workbook()
    wb.remove(wb.active)

    # ---- chargement des sources
    produits = load_produits()
    trans = load_transactions(produits)
    calendar = load_calendar()
    prix = build_prix_keys(trans, calendar)

    # ---- T_Produits : prix_moyen = moyenne des prix des transactions du produit (formule)
    pcols = list(produits.columns)
    c_st, c_it = CL(pcols, "store_id"), CL(pcols, "item_id")
    produits["prix_moyen"] = [
        (f'=IFERROR(AVERAGEIFS(tblTransactions[sell_price],tblTransactions[store_id],${c_st}{DATA_START_ROW + i},'
         f'tblTransactions[item_id],${c_it}{DATA_START_ROW + i}),"")')
        for i in range(len(produits))
    ]
    write_table_sheet(wb, "T_Produits", "Produits suivis",
                      "Clé magasin × produit — prix moyen calculé à partir des transactions.",
                      produits, formats={"prix_moyen": "0.00"})

    p_store, p_cat = (rng("T_Produits", h) for h in ("store_id", "cat_id"))

    cats = sorted(produits["cat_id"].dropna().unique())
    cat_cols = ["cat_id", "description", "nb_produits", "source"]
    cat_rows = []
    for i, c in enumerate(cats):
        r = DATA_START_ROW + i
        cat_rows.append([c, CAT_INFO.get(c, ""),
                         f'=COUNTIFS({p_cat},$B{r},{p_store},sessStoreCrit)', "M5"])
    write_table_sheet(wb, "T_Categories", "Catégories", ", ".join(cats),
                      pd.DataFrame(cat_rows, columns=cat_cols))

    stores = sorted(produits["store_id"].dropna().unique())
    mag_cols = ["store_id", "state_id", "description", "nb_produits", "source"]
    mag_rows = []
    for i, s in enumerate(stores):
        r = DATA_START_ROW + i
        st, desc = STORE_INFO.get(s, (s.split("_")[0], ""))
        mag_rows.append([s, st, desc, f'=COUNTIFS({p_store},$B{r})', "M5"])
    write_table_sheet(wb, "T_Magasins", "Magasins", ", ".join(stores),
                      pd.DataFrame(mag_rows, columns=mag_cols))

    write_table_sheet(wb, "T_Transactions", "Transactions (historique des ventes)",
                      "Lignes M5 en lecture seule ; saisies manuelles = source MANUEL.", trans)

    # ---- T_Prix : sell_price hebdomadaire = formule sur T_Transactions (pas de wm_yr_wk, pas de saisie)
    prix_cols = ["date", "store_id", "item_id", "sell_price"]
    c_d, c_s, c_i = CL(prix_cols, "date"), CL(prix_cols, "store_id"), CL(prix_cols, "item_id")
    prix["sell_price"] = [
        (f'=IFERROR(AVERAGEIFS(tblTransactions[sell_price],tblTransactions[store_id],${c_s}{DATA_START_ROW + i},'
         f'tblTransactions[item_id],${c_i}{DATA_START_ROW + i},'
         f'tblTransactions[date],">="&${c_d}{DATA_START_ROW + i},'
         f'tblTransactions[date],"<="&(${c_d}{DATA_START_ROW + i}+6)),"")')
        for i in range(len(prix))
    ]
    prix = prix[prix_cols]
    write_table_sheet(wb, "T_Prix", "Prix des produits",
                      "Prix hebdomadaires calculés par formules à partir des transactions (semaine = 7 jours depuis la date).",
                      prix, formats={"sell_price": "0.00", "date": "yyyy-mm-dd"})

    # ---- T_Calendrier : weekday / month / year = formules a partir de la date ; wm_yr_wk supprimee
    cal = calendar.drop(columns=["wm_yr_wk"], errors="ignore").copy()
    cal["date"] = cal["date"].dt.date
    for k, name in enumerate(["weekday", "month", "year"]):
        if name not in cal.columns:
            cal.insert(cal.columns.get_loc("date") + 1 + k, name, None)
    ccols = list(cal.columns) + ["source"]
    dl = CL(ccols, "date")
    days = '"Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"'
    n_cal = len(cal)
    cal["weekday"] = [f'=IF(${dl}{DATA_START_ROW + i}="","",CHOOSE(WEEKDAY(${dl}{DATA_START_ROW + i},2),{days}))'
                      for i in range(n_cal)]
    cal["month"] = [f'=IF(${dl}{DATA_START_ROW + i}="","",MONTH(${dl}{DATA_START_ROW + i}))' for i in range(n_cal)]
    cal["year"] = [f'=IF(${dl}{DATA_START_ROW + i}="","",YEAR(${dl}{DATA_START_ROW + i}))' for i in range(n_cal)]
    cal["source"] = "M5"
    write_table_sheet(wb, "T_Calendrier", "Calendrier et événements",
                      "weekday, month, year calculés à partir de la date · événements manuels = source MANUEL.",
                      cal, formats={"month": "0", "year": "0", "date": "yyyy-mm-dd"})

    # ---- feuilles applicatives
    build_gestion(wb)
    build_powerbi(wb)
    build_ml(wb)
    build_prevision(wb)
    build_r_sheets(wb)
    build_session_sheet(wb)
    build_config(wb)
    ws_accueil = wb.create_sheet("Accueil")
    build_accueil(ws_accueil)

    order = ["Accueil", "PowerBI", "ML", "Prevision", "Gestion",
             "T_Produits", "T_Categories", "T_Magasins",
             "T_Transactions", "T_Prix", "T_Calendrier",
             "R_Previsions", "R_Metriques", "R_Resume", "R_Mensuel",
             "R_MetMensuel", "R_Session", "Config"]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=(i - wb.sheetnames.index(name)))

    # une seule feuille selectionnee (sinon Excel ouvre les feuilles "groupees")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == "Accueil")
    wb.active = wb.sheetnames.index("Accueil")

    wb.save(XLSX_PATH)
    print(f"✔ {XLSX_PATH.name} généré ({len(wb.sheetnames)} feuilles)")
    print("\nÉtape suivante : python build_userforms.py  (Windows + Excel)")


if __name__ == "__main__":
    main()